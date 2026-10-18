# 從 E507 xlsm 抽出所有 Data mapping 範例（12 列 × SW code ＋ 輸出名稱），推出每格是哪個 SW、哪一欄（S），存成 node 測試用 fixture
import openpyxl, re, json, sys
from openpyxl.utils import get_column_letter as L
src = sys.argv[1]; out = sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
NM = re.compile(r'^[RGB]-?\d+$')
def lin(n):  # 名稱 → subpixel 線性位置（R1=0）
    c = 'RGB'.index(n[0]); p = int(n[1:]); return (p - 1) * 3 + c if p > 0 else p * 3 + c
def pool(c, v):  # EN01 pool：第 c 欄（I=0）code v ⇒ subpixel；= poolIdx(12, c+1, v)
    return 12 * c - 12 * (v // 12) + v % 12
def name(k):
    p = k // 3; c = 'RGB'[k % 3]; return c + str(p + 1 if k >= 0 else p)
ex = []
for sh in ['FORCE', 'NORMAL', 'ZIGZAG', 'HSD', 'MUX2', 'MUX3', 'MUX4', 'MUX6', 'New BOE', '工作表8']:
    ws = wb[sh]; mr = ws.max_row; mc = min(ws.max_column, 60)
    g = lambda r, c: ws.cell(r, c).value
    r = 1
    while r <= mr - 11:
        found = None
        if any(isinstance(g(r, cc), str) and g(r, cc).strip() in ('mirror', 'line', 'data_sft') for cc in range(1, mc + 1)): r += 1; continue
        for c0 in range(4, 14):
            # 連續整數 code 欄 c0..c1，右邊（可隔一欄）接名稱
            if not isinstance(g(r, c0), (int, float)) or isinstance(g(r, c0), bool): continue
            if c0 > 1 and isinstance(g(r, c0 - 1), (int, float)) and not isinstance(g(r, c0 - 1), bool): continue
            c1 = c0
            while isinstance(g(r, c1 + 1), (int, float)) and not isinstance(g(r, c1 + 1), bool): c1 += 1
            m = c1 - c0 + 1
            if m not in (1, 2, 3, 4, 6): continue
            ok = True
            for rr in range(r, r + 12):
                for cc in range(c0, c1 + 1):
                    v = g(rr, cc)
                    if not isinstance(v, (int, float)) or isinstance(v, bool) or not (0 <= v <= 255): ok = False
            if not ok: continue
            # 輸出名稱在 c1+1 之後
            outs = []
            # 輸出欄範圍：從 c1+1 往右，遇到另一組 code 欄（整數且不是 0 佔位）或連續 2 欄全空就停
            cend = c1
            cc = c1 + 1; empty = 0; started = False
            while cc <= mc:
                col = [g(rr, cc) for rr in range(r, r + 12)]
                hasNm = any(isinstance(v, str) and NM.match(v) for v in col)
                isInt = all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in col) and any(v not in (0,) for v in col)
                if hasNm: started = True; empty = 0; cend = cc
                elif started and isInt: break
                elif all(v is None for v in col):
                    empty += 1
                    if started and empty >= 2: break
                cc += 1
            for rr in range(r, r + 12):
                for cc in range(c1 + 1, cend + 1):
                    v = g(rr, cc)
                    if isinstance(v, str) and NM.match(v): outs.append((rr - r, cc, v))
            if len(outs) >= 12: found = (c0, c1, m, outs); break
        if not found: r += 1; continue
        c0, c1, m, outs = found
        codes = [[int(g(r + i, c0 + j)) for j in range(m)] for i in range(12)]
        if m == 1 and [x[0] for x in codes] == list(range(codes[0][0], codes[0][0] + 12)):
            # pool 表本身（code 連續 12 個、整列名稱）：不是範例；但有 mirror 標籤的（NORMAL 47）照收
            hasMir = any(g(rr, cc) == 'mirror' for rr in range(max(1, r - 4), r) for cc in range(1, mc + 1))
            if not hasMir: r += 12; continue
        mir = None
        for rr in range(max(1, r - 4), r):
            for cc in range(1, mc + 1):
                if g(rr, cc) == 'mirror':
                    for c2 in range(cc + 1, cc + 6):
                        if isinstance(g(rr, c2), (int, float)): mir = int(g(rr, c2)); break
        cells = []; amb = 0; bad = 0
        for (ch, cc, v) in outs:
            c = cc - 9  # pool 第 0 欄＝I 欄（FORCE!I＝8'h01；各分頁的 pool 表 H＝code、I 起）
            js = [j for j in range(m) if codes[ch][j] < 132 and pool(c, codes[ch][j]) == lin(v)]
            if len(js) == 1: cells.append({'ch': ch, 'col': L(cc), 'c': c, 'j': js[0], 'name': v})
            elif len(js) > 1: amb += 1; cells.append({'ch': ch, 'col': L(cc), 'c': c, 'j': js, 'name': v})
            else: bad += 1; cells.append({'ch': ch, 'col': L(cc), 'c': c, 'j': None, 'name': v})
        ex.append({'sheet': sh, 'row0': r, 'codeCols': L(c0) + ':' + L(c1), 'm': m, 'mirror': mir, 'codes': codes, 'cells': cells, 'amb': amb, 'nomatch': bad})
        r += 12
for e in ex:
    js = [c for c in e['cells'] if isinstance(c['j'], int)]
    first = {}
    for c in js: first[c['j']] = min(first.get(c['j'], 99), c['c'])
    sF = first.get(0); m = e['m']; e['sFirst'] = sF; viol = 0
    for c in js:
        d = c['c'] - sF - c['j']
        if d % m or d < 0: viol += 1
        else: c['b'] = d // m
    e['viol'] = viol
json.dump(ex, open(out, 'w'), ensure_ascii=False, indent=0)
print('TOTAL examples', len(ex), 'cells', sum(len(e['cells']) for e in ex), 'ambiguous', sum(e['amb'] for e in ex), 'nomatch', sum(e['nomatch'] for e in ex), 'progression violations', sum(e['viol'] for e in ex))
for e in ex:
    js = [c for c in e['cells'] if isinstance(c['j'], int)]
    # 每個 SW 的最小欄（block 0）與 block 步距
    first = {}
    for c in js: first[c['j']] = min(first.get(c['j'], 99), c['c'])
    print(e['sheet'], e['row0'], e['codeCols'], 'm=', e['m'], 'mir=', e['mirror'], 'cells', len(e['cells']), 'amb', e['amb'], 'nomatch', e['nomatch'], 'SWfirstCol', dict(sorted(first.items())))
