# 產生 EN01 Auto 對照（mirror=0）：datamap-core.js EN01_AUTO_EX 的 JS 片段＋tools/datamap_en01_auto_examples.json（逐格驗證用）
# 用法：python3 tools/datamap_en01_auto_build.py <E507_rt7_sram_w_hand(1T4P).xlsm> tools/datamap_en01_xlsm_examples.json tools/datamap_en01_auto_examples.json <輸出 js 片段>
import openpyxl, re, json, sys
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
fx = json.load(open(sys.argv[2]))
NM = re.compile(r'^[RGB]-?\d+$')
L = openpyxl.utils.get_column_letter
def blk(sh, row, col, m, sF):
    ws = wb[sh]; c0 = openpyxl.utils.column_index_from_string(col)
    codes = [[int(ws.cell(row + i, c0 + j).value) for j in range(m)] for i in range(12)]
    cells = []
    # 輸出區＝從 I 欄往右，遇到連續 3 欄（12 列）都沒有名稱就停（HSD 分頁 AB 欄起另有一份重複的顯示，不算）
    cend = 9; gap = 0
    for c in range(9, 9 + 60):
        if any(isinstance(ws.cell(row + i, c).value, str) and NM.match(ws.cell(row + i, c).value.strip()) for i in range(12)): cend = c; gap = 0
        else:
            gap += 1
            if gap >= 3 and cend > 9: break
    for ch in range(12):
        for c in range(9, cend + 1):
            v = ws.cell(row + ch, c).value
            if not (isinstance(v, str) and NM.match(v.strip())): continue
            k = c - 9 - sF
            if k < 0: continue
            cells.append({'ch': ch, 'j': k % m, 'b': k // m, 'col': L(c), 'name': v.strip()})
    return {'ref': sh + '!' + col + str(row), 'm': m, 'sF': sF, 'codes': codes, 'cells': cells}
# key：kind/panel/sub 條件 → [line0 區塊, line1 區塊]
B = {
 'normal':      [('NORMAL', 47, 'H', 1, 0)],
 'zz_sft0':     [('ZIGZAG', 48, 'H', 1, 1)],
 'zz_sft1':     [('ZIGZAG', 64, 'H', 1, 1)],
 'hsd1':        [('HSD', 148, 'H', 2, 1)],
 'hsd4':        [('HSD', 213, 'H', 2, 1)],
 'hsd35':       [('HSD', 280, 'H', 2, 1), ('HSD', 296, 'H', 2, 1)],
 'hsd4zz':      [('HSD', 348, 'H', 2, 1), ('HSD', 365, 'H', 2, 1)],
 'mux2_n1':     [('MUX2', 97, 'H', 2, 2)],
 'mux2_n2':     [('MUX2', 128, 'H', 2, 2)],
 'mux2_z1':     [('MUX2', 162, 'H', 2, 2), ('MUX2', 178, 'H', 2, 2)],
 'mux2_z2':     [('MUX2', 227, 'H', 2, 2), ('MUX2', 243, 'H', 2, 2)],
 'mux3_n1':     [('MUX3', 97, 'H', 3, 2)],
 'mux3_n2':     [('MUX3', 128, 'H', 3, 2)],
 'mux3_z1':     [('MUX3', 160, 'H', 3, 2), ('MUX3', 176, 'H', 3, 2)],
 'mux3_z2':     [('MUX3', 225, 'H', 3, 2), ('MUX3', 241, 'H', 3, 2)],
 'mux4_n1':     [('MUX4', 107, 'G', 4, 3)],
 'mux4_z1':     [('MUX4', 139, 'G', 4, 3), ('MUX4', 155, 'G', 4, 3)],
 'mux6_n1':     [('MUX6', 143, 'H', 6, 5)],
 'mux6_n2':     [('MUX6', 183, 'H', 6, 5)],
 'mux6_z1':     [('MUX6', 215, 'H', 6, 5), ('MUX6', 231, 'H', 6, 5)],
}
out = {}; js = []
for k, lst in B.items():
    bl = [blk(*x) for x in lst]
    out[k] = bl
    js.append("    %s: { m: %d, sF: %d, ref: %s, L: [%s] }" % (k, bl[0]['m'], bl[0]['sF'], json.dumps([b['ref'] for b in bl]),
              ', '.join('[' + ','.join(('%d' % r[0]) if len(r) == 1 else '[' + ','.join(map(str, r)) + ']' for r in b['codes']) + ']' for b in bl)))
# fixture 對 sF 交叉檢查
for e in fx:
    for k, bl in out.items():
        for b in bl:
            if b['ref'].split('!')[0] == e['sheet'] and int(b['ref'].split('!')[1][1:]) == e['row0']:
                if b['sF'] != e['sFirst'] or b['codes'] != e['codes']: print('MISMATCH', k, b['ref'], b['sF'], e['sFirst'])
json.dump(out, open(sys.argv[3], 'w'), ensure_ascii=False)
open(sys.argv[4], 'w').write(',\n'.join(js) + '\n')
print('ok', len(out), sum(len(b['cells']) for bl in out.values() for b in bl))
